"""Build the five-slide DASO discussion deck and its matching vector PDF.

Run in this directory: python build_deck.py
Dependencies: python-pptx, reportlab. Local font/data assets are bundled.
"""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
DATA = json.loads((ASSETS / "data.json").read_text())
W, H = 960, 540
FONT = "Noto Sans SC"
NAVY = "172B45"
BLUE = "285BA7"
TEAL = "11877E"
MUTED = "5F6F82"
LIGHT = "F3F6FA"
PALE = "EDF6F5"
GRID = "DCE4ED"
WHITE = "FFFFFF"


class Deck:
    """Use the same slide geometry for editable PowerPoint and vector PDF."""

    def __init__(self):
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Pt(W), Pt(H)
        self.ppt.core_properties.title = "DASO：从压缩方法到部署决策"
        self.ppt.core_properties.subject = "技术交流：DASO 方法与实验结果"
        self.ppt.core_properties.author = "DASO"
        self.pdf = canvas.Canvas(str(ROOT / "DASO_技术交流.pdf"), pagesize=(W, H))
        self.pdf.setTitle(self.ppt.core_properties.title)
        self.pdf.setAuthor("DASO")
        for weight in ("Regular", "Bold"):
            pdfmetrics.registerFont(TTFont(f"DASO-{weight}", str(ASSETS / "fonts" / f"NotoSansSC-{weight}.ttf")))
        self.slide = None
        self.text_log = []

    def rect(self, x, y, w, h, fill=WHITE, border=None, radius=0):
        shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = self.slide.shapes.add_shape(shape_type, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if border:
            shape.line.color.rgb = RGBColor.from_string(border)
            shape.line.width = Pt(0.8)
        else:
            shape.line.fill.background()
        if radius:
            shape.adjustments[0] = min(0.12, radius / min(w, h))
        self.pdf.setFillColorRGB(*self.rgb(fill))
        self.pdf.setStrokeColorRGB(*self.rgb(border or fill))
        self.pdf.setLineWidth(0.8)
        self.pdf.roundRect(x, H - y - h, w, h, radius, fill=1, stroke=int(bool(border)))
        return shape

    @staticmethod
    def rgb(color):
        return tuple(int(color[i:i + 2], 16) / 255 for i in (0, 2, 4))

    def line(self, x1, y1, x2, y2, color=GRID, width=1, dashed=False):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.line.color.rgb = RGBColor.from_string(color)
        shape.line.width = Pt(width)
        if dashed:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        self.pdf.setStrokeColorRGB(*self.rgb(color))
        self.pdf.setLineWidth(width)
        self.pdf.setDash([4, 3] if dashed else [])
        self.pdf.line(x1, H - y1, x2, H - y2)
        self.pdf.setDash([])

    def text(self, x, y, w, h, value, size=16, color=NAVY, bold=False, align="left", leading=1.23):
        font_name = "DASO-Bold" if bold else "DASO-Regular"
        lines = value.split("\n")
        for line in lines:
            width = pdfmetrics.stringWidth(line, font_name, size)
            if width > w + 1:
                raise ValueError(f"Text too wide ({width:.1f}>{w}): {line}")
        if size + (len(lines) - 1) * size * leading > h + 2:
            raise ValueError(f"Text too high: {value}")
        self.text_log.append({"page": len(self.ppt.slides), "text": value, "size": size, "box": [x, y, w, h]})
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        frame = box.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = Pt(0)
        frame.margin_top = frame.margin_bottom = Pt(0)
        frame.word_wrap = False
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(lines):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(size * leading)
            run = p.add_run()
            run.text = line
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            prop = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                node = OxmlElement(tag)
                node.set("typeface", FONT)
                prop.append(node)
            self.pdf.setFillColorRGB(*self.rgb(color))
            self.pdf.setFont(font_name, size)
            width = pdfmetrics.stringWidth(line, font_name, size)
            tx = x + (w - width) * {"left": 0, "center": 0.5, "right": 1}[align]
            self.pdf.drawString(tx, H - y - size * 1.02 - index * size * leading, line)
        return box

    def picture(self, x, y, w, h, filename):
        path = ASSETS / "paper_figures" / filename
        with Image.open(path) as image:
            scale = min(w / image.width, h / image.height)
            pw, ph = image.width * scale, image.height * scale
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
        self.slide.shapes.add_picture(str(path), Pt(px), Pt(py), width=Pt(pw), height=Pt(ph))
        self.pdf.drawImage(ImageReader(str(path)), px, H - py - ph, pw, ph, mask="auto")

    def arrow(self, x, y, length=28, color=MUTED):
        self.line(x, y, x + length, y, color, 1.4)
        self.line(x + length - 6, y - 4, x + length, y, color, 1.4)
        self.line(x + length - 6, y + 4, x + length, y, color, 1.4)

    def start(self, chapter, title, subtitle, source, notes):
        self.slide = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        self.rect(0, 0, 8, H, BLUE)
        self.text(46, 24, 820, 18, f"DASO  /  {chapter}", 10.5, TEAL, True)
        self.text(46, 61, 868, 45, title, 29, bold=True)
        self.text(46, 111, 868, 25, subtitle, 13.5, MUTED)
        self.text(46, 488, 868, 17, source, 8.4, MUTED)
        self.line(46, 511, 914, 511)
        self.text(46, 519, 780, 14, "DASO  ·  技术交流", 8.5, MUTED)
        self.text(844, 519, 70, 14, f"{len(self.ppt.slides):02d} / 05", 8.5, MUTED, align="right")
        self.slide.notes_slide.notes_text_frame.text = notes

    def takeaway(self, value, y=432, size=17):
        self.rect(46, y, 868, 45, PALE, radius=5)
        self.rect(46, y, 4, 45, TEAL)
        self.text(64, y + 11, 831, 25, value, size, TEAL, True)

    def end(self):
        self.pdf.showPage()

    def save(self):
        self.ppt.save(ROOT / "DASO_技术交流.pptx")
        self.pdf.save()
        (ROOT / "layout_check.json").write_text(json.dumps({"slides": len(self.ppt.slides), "text_boxes": self.text_log}, ensure_ascii=False, indent=2))


def background(d):
    d.start("背景与问题", "DASO：从压缩方法，到部署决策", "Deployment-Aware Strategy Optimization · 面向部署的压缩策略优化", "依据：最终论文 Introduction / Framework Overview。", "开场：我们研究的是怎样把已有压缩方法转化为实际部署策略。给定模型、硬件、负载和质量要求，选择各模块、各阶段的压缩方法。当前工作来自采用 AAAI 2027 模板的 DASO 论文；此处不宣称录用状态。\n论文正式标题：From Compression Methods to Deployment Decisions: Quality-Constrained Fine-Grained Strategy Optimization for Efficient LLM Inference。\n来源：artifacts/paper/wu.tex。")
    d.text(46, 164, 240, 27, "可选方法越来越多", 18, bold=True)
    methods = ["BF16", "Dense NVFP4", "Sparse BF16", "Sparse NVFP4", "W4A16 Marlin"]
    for i, method in enumerate(methods):
        d.rect(46, 204 + i * 36, 232, 29, LIGHT, radius=4)
        d.text(60, 208 + i * 36, 205, 21, method, 13.5, BLUE)
    d.line(310, 164, 310, 387)
    d.text(347, 164, 560, 27, "实际部署，需要回答三个问题", 18, bold=True)
    items = [
        ("01", "在哪些模块用？", "不同层与模块，对压缩的敏感度不同。"),
        ("02", "Prefill 与 Decode 如何搭配？", "同一种方法，在不同推理阶段的收益不同。"),
        ("03", "质量预算下，能快多少？", "需要面向具体模型、硬件和负载做决策。"),
    ]
    for i, (number, title, body) in enumerate(items):
        y = 208 + i * 67
        d.text(347, y, 30, 23, number, 13, TEAL, True)
        d.text(393, y - 2, 510, 25, title, 17, bold=True)
        d.text(393, y + 24, 510, 22, body, 13, MUTED)
    d.takeaway("研究目标：在质量损失预算内，选择更快的细粒度压缩策略。")
    d.end()


def observation(d):
    d.start("关键观察", "一个方法，难以覆盖所有推理阶段", "RTX 5090 · Llama-2-7B-Chat：均匀 NVFP4 与 W4A16 呈现互补收益", "数据：最终主表；Prefill 为 TTFT 加速，Decode 为 TPOT 加速；BF16 = 1.0×。", "先看动机。Dense NVFP4 的 prefill 加速为 1.71×，decode 为 0.55×；W4A16 Marlin 分别为 1.02×、1.50×。它们说明必须基于实际执行测量判断收益。图中阶段来自论文两种 workload，不能把柱状图数值直接相加当作端到端加速。模块敏感度差异的证据见论文 Quality Modeling。\n来源：daso/results/paper/tables/main_table_final.tex。")
    x0, y0, cw, ch = 90, 374, 440, 178
    d.text(46, 156, 505, 23, "相对 BF16 的阶段加速（×）", 13, MUTED)
    for x, name, color in [(83, "Dense NVFP4", BLUE), (281, "W4A16 Marlin", TEAL)]:
        d.rect(x, 185, 13, 8, color)
        d.text(x + 21, 178, 170, 22, name, 11.5, color)
    for value in [0, 0.5, 1, 1.5, 2]:
        y = y0 - ch * value / 2
        d.line(x0, y, x0 + cw, y, MUTED if value == 1 else GRID, 0.9, dashed=value == 1)
        d.text(46, y - 8, 33, 18, f"{value:.1f}", 10, MUTED, align="right")
    d.text(535, y0 - ch / 2 - 8, 55, 19, "BF16", 10, MUTED)
    main = DATA["main"]
    series = [("Prefill", [main["llama2|Dense NVFP4"][4], main["llama2|W4A16 Marlin"][4]]),
              ("Decode", [main["llama2|Dense NVFP4"][9], main["llama2|W4A16 Marlin"][9]])]
    for i, (label, values) in enumerate(series):
        center = x0 + 118 + i * 208
        for j, (value, color) in enumerate(zip(values, [BLUE, TEAL])):
            x = center - 54 + j * 62
            bh = ch * value / 2
            d.rect(x, y0 - bh, 44, bh, color)
            d.text(x - 9, y0 - bh - 28, 62, 23, f"{value:.2f}×", 13, color, True, "center")
        d.text(center - 70, y0 + 13, 140, 25, label, 14, bold=True, align="center")
    d.line(616, 168, 616, 403)
    d.text(649, 183, 265, 27, "阶段差异", 18, BLUE, True)
    d.text(649, 219, 265, 57, "需要分别选择\nPrefill 与 Decode 的方法。", 14, MUTED, leading=1.5)
    d.text(649, 302, 265, 27, "模块差异", 18, TEAL, True)
    d.text(649, 338, 265, 58, "保留模块之间的敏感度差异，\n避免对所有模块统一压缩。", 13.5, MUTED, leading=1.5)
    d.takeaway("启发：让各模块、各推理阶段分别选择合适的方法。")
    d.end()


def modeling(d):
    d.start("方法（一）", "建立可比较的时延与质量代价", "面向每个“模块 × 推理阶段 × 候选方法”，估计速度收益与质量损失", "依据：最终论文 Latency / Quality Modeling；右下为论文质量建模原图。", "这一页回答候选方法的代价如何估计。时延模型使用方法感知的 Roofline 先验，区分计算、访存、量化处理与固定开销，再通过实测形状的全局与局部残差校准。质量模型从缓存激活计算局部输出扰动，并以少量完整策略的 teacher-forced NLL 标签拟合方法、层深、模块类型与推理阶段的共享敏感度。局部误差本身不足以预测整体质量，原图展示层深与模块类型造成的敏感度差异。优化使用相对 BF16 的 ΔNLL，最终仍以真实下游任务指标检验。\n来源：artifacts/paper/wu.tex，Method；daso/results/paper/figures/quality_modeling_insight.png。")
    d.text(46, 159, 416, 29, "时延建模：硬件与形状感知", 18, BLUE, True)
    d.text(46, 206, 416, 25, "方法感知的 Roofline 先验", 15.5, bold=True)
    d.text(46, 239, 416, 47, "按方法区分计算、访存和固定开销，\n建立可跨形状比较的时延基线。", 13.8, MUTED, leading=1.5)
    d.text(46, 295, 416, 25, "用实测数据修正预测偏差", 15.5, bold=True)
    d.text(46, 328, 416, 47, "结合全局趋势与邻近形状，\n修正真实内核效率造成的偏差。", 13.8, MUTED, leading=1.5)
    d.rect(46, 389, 416, 35, LIGHT, radius=4)
    d.text(58, 398, 392, 21, "输出：模块 × 阶段 × 方法的时延估计", 12.5, BLUE, True)
    d.line(480, 159, 480, 424)
    d.text(498, 159, 416, 29, "质量建模：从局部扰动到 ΔNLL", 18, TEAL, True)
    d.text(498, 206, 416, 25, "局部输出扰动 + 少量策略级 NLL", 15.2, bold=True)
    d.text(498, 239, 416, 43, "按方法、层深、模块类型与阶段校准，\n将局部误差映射为整体质量损失。", 13.3, MUTED, leading=1.5)
    d.picture(498, 278, 416, 147, "quality_modeling_insight.png")
    d.takeaway("输出统一的时延—质量代价，支持跨方法的策略比较。")
    d.end()


def optimization(d):
    d.start("方法（二）", "预算约束优化与实际推理执行", "将预测代价转化为模块级、阶段级策略，并落实到 vLLM 运行时", "依据：最终论文 Strategy Optimization and Deployment；右侧时序为解释性示意。", "这一页回答代价如何变成可执行策略。每个模块的动作是一对 prefill / decode 方法。时延按 prefill 一次及 decode 的 T-1 次执行汇总，质量成本相加并保留回归截距。给定质量预算后，问题转化为多选背包，用动态规划求解；扫描预算得到多个折中策略。运行时扩展 vLLM 的阶段感知后端分派，各层完成 prefill 后，在独立 copy stream 准备其 decode 权重表示，并与后续层的 prefill 计算重叠。示意图只表达依赖和重叠关系，不表达实际时间比例。\n来源：artifacts/paper/wu.tex，Strategy Optimization and Deployment。")
    d.text(46, 159, 416, 29, "策略优化：在质量预算内最小化时延", 18, BLUE, True)
    d.rect(46, 205, 416, 40, PALE, radius=4)
    d.text(58, 215, 392, 24, "每个模块选择一对 Prefill / Decode 方法", 13.9, TEAL, True)
    for i, text in enumerate(["1  按执行次数汇总阶段时延", "2  以 ΔNLL 限制总质量损失", "3  动态规划选择各模块的方法对"]):
        d.text(46, 271+i*35, 416, 25, text, 14.2, MUTED)
    d.rect(46, 384, 416, 40, LIGHT, radius=4)
    d.text(58, 395, 392, 24, "扫描质量预算 → 生成不同折中策略", 13.3, BLUE, True)
    d.line(480, 159, 480, 424)
    d.text(498, 159, 416, 29, "运行时：让异构策略真正跑起来", 18, TEAL, True)
    d.text(498, 206, 416, 25, "按模块与推理阶段分派不同计算后端。", 13.8, MUTED)
    d.text(498, 254, 416, 23, "逐层计算与 Decode 权重准备重叠", 12.5, bold=True)
    d.text(498, 295, 65, 22, "Prefill", 11.8, BLUE)
    d.text(498, 342, 65, 22, "准备", 11.8, TEAL)
    for i in range(4):
        x=567+i*67
        d.rect(x, 291, 59, 28, BLUE, radius=3)
        d.text(x, 296, 59, 20, f"L{i+1}", 11, WHITE, True, "center")
        d.line(x+59, 319, x+67, 338, GRID, 0.8, dashed=True)
        d.rect(x+67, 338, 59, 28, TEAL, radius=3)
        d.text(x+67, 343, 59, 20, f"L{i+1}", 11, WHITE, True, "center")
    d.line(827, 305, 888, 305, BLUE, 1.2)
    d.arrow(888, 305, 26, BLUE)
    d.arrow(895, 352, 19, TEAL)
    d.text(498, 375, 416, 18, "时序示意；方块不表示实际持续时间。", 9.5, MUTED)
    d.text(498, 404, 416, 24, "通过异步预取，减少阶段切换的全局等待。", 13.4, MUTED)
    d.takeaway("质量预算决定策略，运行时将策略转化为实际推理加速。")
    d.end()


def results(d):
    d.start("实验结果", "实验结果：速度收益与质量折中", "原图：Llama-3.1-8B-Instruct · RTX 5090；右侧摘要来自最终主表", "左：论文质量—速度原图；右：最终主表（两模型，RTX 5090）。B=8，S=2048，生成长度 T=64。", "本页直接嵌入论文最终的双面板原图，不重新绘制。左图是 Llama-3.1-8B-Instruct 的 prefill-only ARC-Challenge acc_norm，右图是 prefill-decode IWSLT SacreBLEU；黑点是实测前沿，红方块是 uniform baselines，蓝菱形为 Balanced，金色星形为 Max speed。图内结果均来自 Llama-3.1。旁边数值摘要来自最终主表的两模型结果：2.26× prefill 来自 Llama-3.1，1.55× E2E 来自 Llama-2；两者是宽松质量预算下的端点，不表示无损。Balanced 的 1.14–1.26× E2E 分别来自 Llama-2 与 Llama-3.1。原图与主表都是 RTX 5090，正文还提供 RTX PRO 6000 的跨平台验证。\n原图：daso/results/paper/figures/llama31_two_scenario_tradeoff_horizontal.png，按文件原样嵌入。\n数值：daso/results/paper/tables/main_table_final.tex。")
    d.picture(46, 150, 550, 327, "llama31_two_scenario_tradeoff_horizontal.png")
    d.line(612, 158, 612, 474)
    d.text(638, 161, 276, 26, "论文中的主要结果", 17.5, bold=True)
    main=DATA["main"]
    prefill=max(row[4] for row in main.values())
    e2e=max(row[10] for row in main.values())
    balanced=[main[f"{model}|DASO (Balanced)"][10] for model in ("llama2", "llama31")]
    d.text(638, 207, 276, 43, f"{prefill:.2f}×", 32, BLUE, True)
    d.text(638, 253, 276, 23, "最大 Prefill 加速 · Llama-3.1", 11.7, MUTED)
    d.text(638, 296, 276, 43, f"{e2e:.2f}×", 32, BLUE, True)
    d.text(638, 342, 276, 23, "最大 E2E 加速 · Llama-2", 11.7, MUTED)
    d.line(638, 375, 914, 375)
    d.text(638, 389, 276, 34, f"{min(balanced):.2f}–{max(balanced):.2f}×", 25, TEAL, True)
    d.text(638, 427, 276, 22, "Balanced 的 E2E 加速", 12.3, MUTED)
    d.text(638, 453, 276, 29, "最大加速使用宽松质量预算；\nBalanced 体现较保守的质量折中。", 10.2, MUTED, leading=1.35)
    d.end()


def main():
    deck = Deck()
    for build in (background, observation, modeling, optimization, results):
        build(deck)
    deck.save()
    print("Generated 5-slide PPTX and PDF in", ROOT)


if __name__ == "__main__":
    main()
